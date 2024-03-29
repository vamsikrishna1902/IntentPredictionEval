from __future__ import division
import sys
import os
import time
import QueryExecution as QExec
from bitmap import BitMap
import CFCosineSim
import TupleIntent as ti
import ParseConfigFile as parseConfig
from ParseConfigFile import getConfig
import pickle
import argparse
from pandas import DataFrame
from openpyxl import load_workbook
import pandas as pd

def updateArrWithDictEntry(arr, evalOpsObjDict, epIndex, numOpQueryCountDict):
    try:
        arr.append(float(evalOpsObjDict[epIndex])/float(numOpQueryCountDict[epIndex]))
    except:
        arr.append("")
    return

def updateArrWithCountEntry(arr, numOpQueryCountDict, key):
    try:
        arr.append(numOpQueryCountDict[key])
    except:
        arr.append(0.0)
    return

def updateAggMetricWithDictEntry(avgMetric, evalOpsObjDict, epIndex):
    try:
        avgMetric += float(evalOpsObjDict[epIndex])
    except:
        avgMetric += 0.0
    return avgMetric

def plotQueryTypeDistribution(evalOpsObj):
    episodes = []
    numSelectQueryType = []
    numInsertQueryType = []
    numUpdateQueryType = []
    numDeleteQueryType = []
    totalSelect = 0.0
    totalInsert = 0.0
    totalUpdate = 0.0
    totalDelete = 0.0
    for key in sorted(evalOpsObj.meanReciprocalRank.keys()):
        episodes.append(key)
        updateArrWithCountEntry(numSelectQueryType, evalOpsObj.numSelectQueryType, key)
        updateArrWithCountEntry(numInsertQueryType, evalOpsObj.numInsertQueryType, key)
        updateArrWithCountEntry(numUpdateQueryType, evalOpsObj.numUpdateQueryType, key)
        updateArrWithCountEntry(numDeleteQueryType, evalOpsObj.numDeleteQueryType, key)
        totalSelect = updateAggMetricWithDictEntry(totalSelect, evalOpsObj.numSelectQueryType, key)
        totalInsert = updateAggMetricWithDictEntry(totalInsert, evalOpsObj.numInsertQueryType, key)
        totalUpdate = updateAggMetricWithDictEntry(totalUpdate, evalOpsObj.numUpdateQueryType, key)
        totalDelete = updateAggMetricWithDictEntry(totalDelete, evalOpsObj.numDeleteQueryType, key)
    df = DataFrame(
        {'episodes': episodes, '# SELECT': numSelectQueryType, '# INSERT': numInsertQueryType, '# UPDATE': numUpdateQueryType,
         '# DELETE': numDeleteQueryType})
    algoName = evalOpsObj.configDict['ALGORITHM']
    if evalOpsObj.configDict['ALGORITHM'] == 'RNN' and evalOpsObj.configDict['RNN_PREDICT_NOVEL_QUERIES'] == 'True':
        algoName = "NovelRNN"
    elif evalOpsObj.configDict['ALGORITHM'] == 'RNN' and evalOpsObj.configDict['RNN_PREDICT_NOVEL_QUERIES'] == 'False':
        algoName = "HistoricalRNN"
    outputOpWiseQualityFileName = getConfig(evalOpsObj.configDict['OUTPUT_DIR']) + "/OpWiseExcel/QTDist" + \
                                  algoName
    df.to_excel(outputOpWiseQualityFileName + ".xlsx", sheet_name='sheet1', index=False)
    totalSelList = []
    totalSelList.append(totalSelect)
    totalInsList = []
    totalInsList.append(totalInsert)
    totalUpdList = []
    totalUpdList.append(totalUpdate)
    totalDelList = []
    totalDelList.append(totalDelete)
    df = DataFrame({'totalSEL': totalSelList, 'totalINS': totalInsList, 'totalUPD': totalUpdList, 'totalDEL': totalDelList})
    outputOpWiseQualityFileName = getConfig(evalOpsObj.configDict['OUTPUT_DIR']) + "/OpWiseExcel/TotalQT_" + \
                                  algoName
    df.to_excel(outputOpWiseQualityFileName + ".xlsx", sheet_name='sheet2', index=False)
    return


def plotMeanReciprocalRank(evalOpsObj):
    episodes = []
    meanReciprocalRank = []
    numEpQueries = []
    avgMRR = 0.0
    for key in sorted(evalOpsObj.meanReciprocalRank.keys()):
        episodes.append(key)
        updateArrWithCountEntry(numEpQueries, evalOpsObj.numEpQueries, key)
        updateArrWithDictEntry(meanReciprocalRank, evalOpsObj.meanReciprocalRank, key, evalOpsObj.numEpQueries)
        avgMRR = updateAggMetricWithDictEntry(avgMRR, evalOpsObj.meanReciprocalRank, key)
    df = DataFrame(
        {'episodes': episodes, 'meanReciprocalRank': meanReciprocalRank, 'numMRRQueries': numEpQueries})
    algoName = evalOpsObj.configDict['ALGORITHM']
    if evalOpsObj.configDict['ALGORITHM'] == 'RNN' and evalOpsObj.configDict['RNN_PREDICT_NOVEL_QUERIES'] == 'True':
        algoName = "NovelRNN"
    elif evalOpsObj.configDict['ALGORITHM'] == 'RNN' and evalOpsObj.configDict['RNN_PREDICT_NOVEL_QUERIES'] == 'False':
        algoName = "HistoricalRNN"
    outputOpWiseQualityFileName = getConfig(evalOpsObj.configDict['OUTPUT_DIR']) + "/OpWiseExcel/Output_MRR_" + algoName
    df.to_excel(outputOpWiseQualityFileName + ".xlsx", sheet_name='sheet1', index=False)
    totalQueryCount = float(sum(numEpQueries))
    if totalQueryCount > 0.0:
        avgMRR = avgMRR / totalQueryCount
    avgMRRList = []
    avgMRRList.append(avgMRR)
    totalQueryCountList = []
    totalQueryCountList.append(totalQueryCount)
    df = DataFrame({'avgMRR': avgMRRList, 'numMRRQueries': totalQueryCountList})
    outputOpWiseQualityFileName = getConfig(evalOpsObj.configDict['OUTPUT_DIR']) + "/OpWiseExcel/AggrOutput_MRR_" + algoName
    df.to_excel(outputOpWiseQualityFileName + ".xlsx", sheet_name='sheet2', index=False)
    return

def plotOp(evalOpsP, evalOpsR, evalOpsF, numOpQueryCountDict, evalOpsObj, opString):
    episodes = []
    resP = []
    resR = []
    resF = []
    numEpQueries = []
    avgResP = 0.0
    avgResR = 0.0
    avgResF = 0.0
    for key in sorted(evalOpsObj.queryTypeP.keys()):
        episodes.append(key)
        updateArrWithCountEntry(numEpQueries, numOpQueryCountDict, key)
        updateArrWithDictEntry(resP, evalOpsP, key, numOpQueryCountDict)
        updateArrWithDictEntry(resR, evalOpsR, key, numOpQueryCountDict)
        updateArrWithDictEntry(resF, evalOpsF, key, numOpQueryCountDict)
        avgResP = updateAggMetricWithDictEntry(avgResP, evalOpsP, key)
        avgResR = updateAggMetricWithDictEntry(avgResR, evalOpsR, key)
        avgResF = updateAggMetricWithDictEntry(avgResF, evalOpsF, key)
    headerP = evalOpsObj.configDict['ALGORITHM']+"(P)"
    headerR = evalOpsObj.configDict['ALGORITHM']+"(R)"
    headerF = evalOpsObj.configDict['ALGORITHM'] + "(F)"
    headerQ = 'num'+opString+'Queries'
    df = DataFrame(
        {'episodes': episodes, headerP: resP, headerR: resR, headerF: resF, headerQ:numEpQueries})
    algoName = evalOpsObj.configDict['ALGORITHM']
    if evalOpsObj.configDict['ALGORITHM'] == 'RNN' and evalOpsObj.configDict['RNN_PREDICT_NOVEL_QUERIES'] == 'True':
        algoName = "NovelRNN"
    elif evalOpsObj.configDict['ALGORITHM'] == 'RNN' and evalOpsObj.configDict['RNN_PREDICT_NOVEL_QUERIES'] == 'False':
        algoName = "HistoricalRNN"
    outputOpWiseQualityFileName = getConfig(evalOpsObj.configDict['OUTPUT_DIR']) + "/OpWiseExcel/Output_" + opString + "_" + \
                                  algoName
    df.to_excel(outputOpWiseQualityFileName + ".xlsx", sheet_name='sheet1', index=False)
    totalQueryCount = float(sum(numEpQueries))
    if totalQueryCount > 0.0:
        avgResP = avgResP / totalQueryCount
        avgResR = avgResR / totalQueryCount
        avgResF = avgResF / totalQueryCount
    avgResPList = []
    avgResPList.append(avgResP)
    avgResRList = []
    avgResRList.append(avgResR)
    avgResFList = []
    avgResFList.append(avgResF)
    totalQueryCountList = []
    totalQueryCountList.append(totalQueryCount)
    df = DataFrame({headerP: avgResPList, headerR: avgResRList, headerF: avgResFList, headerQ: totalQueryCountList})
    outputOpWiseQualityFileName = getConfig(
        evalOpsObj.configDict['OUTPUT_DIR']) + "/OpWiseExcel/AggrOutput_" + opString + "_" + algoName
    df.to_excel(outputOpWiseQualityFileName + ".xlsx", sheet_name='sheet1', index=False)
    return

def plotEvalMetricsOpWise(evalOpsObj):
    plotQueryTypeDistribution(evalOpsObj)
    plotMeanReciprocalRank(evalOpsObj)
    plotOp(evalOpsObj.queryTypeP, evalOpsObj.queryTypeR, evalOpsObj.queryTypeF, evalOpsObj.numQueryTypeQueries, evalOpsObj, "QUERYTYPE")
    plotOp(evalOpsObj.tablesP, evalOpsObj.tablesR, evalOpsObj.tablesF, evalOpsObj.numTablesQueries, evalOpsObj, "TABLES")
    plotOp(evalOpsObj.projColsP, evalOpsObj.projColsR, evalOpsObj.projColsF, evalOpsObj.numProjColsQueries, evalOpsObj,
           "PROJ")
    plotOp(evalOpsObj.avgColsP, evalOpsObj.avgColsR, evalOpsObj.avgColsF, evalOpsObj.numAvgColsQueries, evalOpsObj,
           "AVG")
    plotOp(evalOpsObj.minColsP, evalOpsObj.minColsR, evalOpsObj.minColsF, evalOpsObj.numMinColsQueries, evalOpsObj,
           "MIN")
    plotOp(evalOpsObj.maxColsP, evalOpsObj.maxColsR, evalOpsObj.maxColsF, evalOpsObj.numMaxColsQueries, evalOpsObj,
           "MAX")
    plotOp(evalOpsObj.sumColsP, evalOpsObj.sumColsR, evalOpsObj.sumColsF, evalOpsObj.numSumColsQueries, evalOpsObj,
           "SUM")
    plotOp(evalOpsObj.countColsP, evalOpsObj.countColsR, evalOpsObj.countColsF, evalOpsObj.numCountColsQueries, evalOpsObj,
           "COUNT")
    plotOp(evalOpsObj.selColsP, evalOpsObj.selColsR, evalOpsObj.selColsF, evalOpsObj.numSelColsQueries, evalOpsObj,
           "SEL")
    plotOp(evalOpsObj.condSelColsP, evalOpsObj.condSelColsR, evalOpsObj.condSelColsF, evalOpsObj.numCondSelColsQueries, evalOpsObj,
           "CONDSEL")
    plotOp(evalOpsObj.groupByColsP, evalOpsObj.groupByColsR, evalOpsObj.groupByColsF, evalOpsObj.numGroupByColsQueries, evalOpsObj,
           "GROUPBY")
    plotOp(evalOpsObj.orderByColsP, evalOpsObj.orderByColsR, evalOpsObj.orderByColsF, evalOpsObj.numOrderByColsQueries, evalOpsObj,
           "ORDERBY")
    plotOp(evalOpsObj.havingColsP, evalOpsObj.havingColsR, evalOpsObj.havingColsF, evalOpsObj.numHavingColsQueries, evalOpsObj,
           "HAVING")
    plotOp(evalOpsObj.limitP, evalOpsObj.limitR, evalOpsObj.limitF, evalOpsObj.numLimitColsQueries, evalOpsObj,
           "LIMIT")
    plotOp(evalOpsObj.joinPredsP, evalOpsObj.joinPredsR, evalOpsObj.joinPredsF, evalOpsObj.numJoinPredsColsQueries, evalOpsObj,
           "JOIN")
    plotOp(evalOpsObj.condJoinPredsP, evalOpsObj.condJoinPredsR, evalOpsObj.condJoinPredsF, evalOpsObj.numCondJoinPredsColsQueries,
           evalOpsObj,
           "CONDJOIN")
    return

class evalOps:
    def __init__(self, configFileName, logFile):
        self.configFileName = configFileName
        self.configDict = parseConfig.parseConfigFile(configFileName)
        self.logFile = logFile
        self.curEpisode = 0
        self.numSelectQueryType = {}
        self.numInsertQueryType = {}
        self.numUpdateQueryType = {}
        self.numDeleteQueryType = {}
        self.numEpQueries = {}
        self.curQueryIndex = -1
        self.meanReciprocalRank = {}
        self.episode = {}
        self.queryTypeP = {}
        self.queryTypeR = {}
        self.queryTypeF = {}
        self.numQueryTypeQueries = {}
        self.tablesP = {}
        self.tablesR = {}
        self.tablesF = {}
        self.numTablesQueries = {}
        self.projColsP = {}
        self.projColsR = {}
        self.projColsF = {}
        self.numProjColsQueries = {}
        self.avgColsP = {}
        self.avgColsR = {}
        self.avgColsF = {}
        self.numAvgColsQueries = {}
        self.minColsP = {}
        self.minColsR = {}
        self.minColsF = {}
        self.numMinColsQueries = {}
        self.maxColsP = {}
        self.maxColsR = {}
        self.maxColsF = {}
        self.numMaxColsQueries = {}
        self.sumColsP = {}
        self.sumColsR = {}
        self.sumColsF = {}
        self.numSumColsQueries = {}
        self.countColsP = {}
        self.countColsR = {}
        self.countColsF = {}
        self.numCountColsQueries = {}
        self.selColsP = {}
        self.selColsR = {}
        self.selColsF = {}
        self.numSelColsQueries = {}
        self.condSelColsP = {}
        self.condSelColsR = {}
        self.condSelColsF = {}
        self.numCondSelColsQueries = {}
        self.groupByColsP = {}
        self.groupByColsR = {}
        self.groupByColsF = {}
        self.numGroupByColsQueries = {}
        self.orderByColsP = {}
        self.orderByColsR = {}
        self.orderByColsF = {}
        self.numOrderByColsQueries = {}
        self.havingColsP = {}
        self.havingColsR = {}
        self.havingColsF = {}
        self.numHavingColsQueries = {}
        self.limitP = {}
        self.limitR = {}
        self.limitF = {}
        self.numLimitColsQueries = {}
        self.joinPredsP = {}
        self.joinPredsR = {}
        self.joinPredsF = {}
        self.numJoinPredsColsQueries = {}
        self.condJoinPredsP = {}
        self.condJoinPredsR = {}
        self.condJoinPredsF = {}
        self.numCondJoinPredsColsQueries = {}

class nextActualOps:
    def __init__(self):
        self.queryType = None
        self.tables = None
        self.projCols = None
        self.avgCols = None
        self.minCols = None
        self.maxCols = None
        self.sumCols = None
        self.countCols = None
        self.selCols = None
        self.groupByCols = None
        self.orderByCols = None
        self.havingCols = None
        self.limit = None
        self.joinPreds = None

def parseLineAddOp(line, actualOrPredObj):
    if line.startswith("Query Type"):
        actualOrPredObj.queryType = line.strip().split(": ")[1]
    elif line.startswith("Limit"):
        actualOrPredObj.limit = line.strip().split(": ")[1]
    elif line.startswith("Tables"):
        actualOrPredObj.tables = eval(line.strip().split(": ")[1])
    elif line.startswith("Projected"):
        actualOrPredObj.projCols = eval(line.strip().split(": ")[1])
    elif line.startswith("AVG"):
        actualOrPredObj.avgCols = eval(line.strip().split(": ")[1])
    elif line.startswith("MIN"):
        actualOrPredObj.minCols = eval(line.strip().split(": ")[1])
    elif line.startswith("MAX"):
        actualOrPredObj.maxCols = eval(line.strip().split(": ")[1])
    elif line.startswith("SUM"):
        actualOrPredObj.sumCols = eval(line.strip().split(": ")[1])
    elif line.startswith("COUNT"):
        actualOrPredObj.countCols = eval(line.strip().split(": ")[1])
    elif line.startswith("SEL"):
        actualOrPredObj.selCols = eval(line.strip().split(": ")[1])
    elif line.startswith("GROUP"):
        actualOrPredObj.groupByCols = eval(line.strip().split(": ")[1])
    elif line.startswith("ORDER"):
        actualOrPredObj.orderByCols = eval(line.strip().split(": ")[1])
    elif line.startswith("HAVING"):
        actualOrPredObj.havingCols = eval(line.strip().split(": ")[1])
    elif line.startswith("JOIN"):
        actualOrPredObj.joinPreds = eval(line.strip().split(": ")[1])
    return

def updateMetricDict(metricDict, key, val):
    if key not in metricDict:
        metricDict[key] = val
    else:
        metricDict[key] = float(metricDict[key]+val)
    return

def computeOpF1(predOpList, actualOpList):
    if predOpList is None and actualOpList is not None:
        return (0.0, 0.0, 0.0)
    elif predOpList is not None and actualOpList is not None:
        TP = len(set(predOpList).intersection(set(actualOpList)))
        FP = len(set(predOpList) - set(actualOpList))
        FN = len(set(actualOpList) - set(predOpList))
        P = float(TP)/float(TP+FP)
        R = float(TP)/float(TP+FN)
        if P == 0.0 or R == 0.0:
            F = 0.0
        else:
            F = 2*P*R / float(P+R)
        return (P, R, F)
    else:
        return (None, None, None)

def updateOpMetrics(P, R, F, evalOpsP, evalOpsR, evalOpsF, evalOpsQueryCountDict, evalOpsObj):
    if P is not None and R is not None and F is not None:
        if evalOpsObj.curEpisode not in evalOpsQueryCountDict:
            evalOpsQueryCountDict[evalOpsObj.curEpisode] = 1
        else:
            evalOpsQueryCountDict[evalOpsObj.curEpisode] += 1
        updateMetricDict(evalOpsP, evalOpsObj.curEpisode, P)
        updateMetricDict(evalOpsR, evalOpsObj.curEpisode, R)
        updateMetricDict(evalOpsF, evalOpsObj.curEpisode, F)
    return

def computeRelevantJoinPreds(accTables, predictedOrActualJoinPreds):
    if predictedOrActualJoinPreds is None:
        return None
    relJoinPreds = []
    for joinPred in predictedOrActualJoinPreds:
        leftTable = joinPred.split(",")[0].split(".")[0]
        rightTable = joinPred.split(",")[1].split(".")[0]
        if leftTable in accTables and rightTable in accTables:
            relJoinPreds.append(joinPred)
    if len(relJoinPreds) == 0:
        return None
    return relJoinPreds

def computeRelevantSelCols(accTables, predOrActualCols):
    if predOrActualCols is None:
        return None
    relCols = []
    for col in predOrActualCols:
        tableName = col.split(".")[0]
        if tableName in accTables:
            relCols.append(col)
    if len(relCols) == 0:
        return None
    return relCols

def compUpdateOpMetrics(predOpList, actualOpList, evalOpsP, evalOpsR, evalOpsF, evalOpsQueryCountDict, evalOpsObj):
    (P,R,F) = computeOpF1(predOpList, actualOpList)
    updateOpMetrics(P, R, F, evalOpsP, evalOpsR, evalOpsF, evalOpsQueryCountDict, evalOpsObj)
    return

def compUpdateCondSelMetrics(predOpsObj, nextActualOpsObj, evalOpsObj):
    try:
        if evalOpsObj.tablesF[evalOpsObj.curEpisode] == 1.0 and evalOpsObj.curEpisode in evalOpsObj.selColsP \
                and evalOpsObj.curEpisode in evalOpsObj.selColsR and evalOpsObj.curEpisode in evalOpsObj.selColsF:
            updateOpMetrics(evalOpsObj.selColsP[evalOpsObj.curEpisode], evalOpsObj.selColsR[evalOpsObj.curEpisode],
                            evalOpsObj.selColsF[evalOpsObj.curEpisode], evalOpsObj.condSelColsP, evalOpsObj.condSelColsR,
                            evalOpsObj.condSelColsF, evalOpsObj.numCondSelColsQueries, evalOpsObj)
        elif evalOpsObj.tablesF[evalOpsObj.curEpisode] > 0.0 and evalOpsObj.curEpisode in evalOpsObj.selColsP \
                and evalOpsObj.curEpisode in evalOpsObj.selColsR and evalOpsObj.curEpisode in evalOpsObj.selColsF: # partial overlap of tables
            accTables = list(set(predOpsObj.tables).intersection(set(nextActualOpsObj.tables)))
            if len(accTables) > 0:
                relPredCols = computeRelevantSelCols(accTables, predOpsObj.selCols)
                relActualCols = computeRelevantSelCols(accTables, nextActualOpsObj.selCols)
                compUpdateOpMetrics(relPredCols, relActualCols, evalOpsObj.condSelColsP, evalOpsObj.condSelColsR,
                                    evalOpsObj.condSelColsF, evalOpsObj.numCondSelColsQueries, evalOpsObj)
    except:
        pass
    return

def compUpdateCondJoinMetrics(predOpsObj, nextActualOpsObj, evalOpsObj):
    try:
        if evalOpsObj.tablesF[evalOpsObj.curEpisode] == 1.0 and evalOpsObj.curEpisode in evalOpsObj.joinPredsP \
                and evalOpsObj.curEpisode in evalOpsObj.joinPredsR and evalOpsObj.curEpisode in evalOpsObj.joinPredsF:
            updateOpMetrics(evalOpsObj.joinPredsP[evalOpsObj.curEpisode], evalOpsObj.joinPredsR[evalOpsObj.curEpisode],
                            evalOpsObj.joinPredsF[evalOpsObj.curEpisode], evalOpsObj.condJoinPredsP, evalOpsObj.condJoinPredsR,
                            evalOpsObj.condJoinPredsF, evalOpsObj.numCondJoinPredsColsQueries, evalOpsObj)
        elif evalOpsObj.tablesF[evalOpsObj.curEpisode] > 0.0 and evalOpsObj.curEpisode in evalOpsObj.joinPredsP \
                and evalOpsObj.curEpisode in evalOpsObj.joinPredsR and evalOpsObj.curEpisode in evalOpsObj.joinPredsF: # partial overlap of tables
            accTables = list(set(predOpsObj.tables).intersection(set(nextActualOpsObj.tables)))
            if len(accTables) > 0:
                relPredictedJoinPreds = computeRelevantJoinPreds(accTables, predOpsObj.joinPreds)
                relActualJoinPreds = computeRelevantJoinPreds(accTables, nextActualOpsObj.joinPreds)
                compUpdateOpMetrics(relPredictedJoinPreds, relActualJoinPreds, evalOpsObj.condJoinPredsP, evalOpsObj.condJoinPredsR,
                                    evalOpsObj.condJoinPredsF, evalOpsObj.numCondJoinPredsQueries, evalOpsObj)
    except:
        pass
    return

def computeF1(evalOpsObj, predOpsObj, nextActualOpsObj):
    if nextActualOpsObj.queryType == "select":
        updateMetricDict(evalOpsObj.numSelectQueryType, evalOpsObj.curEpisode, 1.0)
    elif nextActualOpsObj.queryType == "insert":
        updateMetricDict(evalOpsObj.numInsertQueryType, evalOpsObj.curEpisode, 1.0)
    elif nextActualOpsObj.queryType == "update":
        updateMetricDict(evalOpsObj.numUpdateQueryType, evalOpsObj.curEpisode, 1.0)
    elif nextActualOpsObj.queryType == "delete":
        updateMetricDict(evalOpsObj.numDeleteQueryType, evalOpsObj.curEpisode, 1.0)
    if predOpsObj.queryType == nextActualOpsObj.queryType:
        updateOpMetrics(1.0, 1.0, 1.0, evalOpsObj.queryTypeP, evalOpsObj.queryTypeR,
                        evalOpsObj.queryTypeF, evalOpsObj.numQueryTypeQueries, evalOpsObj)
    elif predOpsObj.queryType != nextActualOpsObj.queryType:
        updateOpMetrics(0.0, 0.0, 0.0, evalOpsObj.queryTypeP, evalOpsObj.queryTypeR,
                        evalOpsObj.queryTypeF, evalOpsObj.numQueryTypeQueries, evalOpsObj)
    if nextActualOpsObj.limit is not None:
        if predOpsObj.limit == nextActualOpsObj.limit:
            updateOpMetrics(1.0, 1.0, 1.0, evalOpsObj.limitP, evalOpsObj.limitR,
                            evalOpsObj.limitF, evalOpsObj.numLimitColsQueries, evalOpsObj)
        else:
            updateOpMetrics(0.0, 0.0, 0.0, evalOpsObj.limitP, evalOpsObj.limitR,
                            evalOpsObj.limitF, evalOpsObj.numLimitColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.tables, nextActualOpsObj.tables, evalOpsObj.tablesP,
                        evalOpsObj.tablesR, evalOpsObj.tablesF, evalOpsObj.numTablesQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.projCols, nextActualOpsObj.projCols, evalOpsObj.projColsP,
                        evalOpsObj.projColsR, evalOpsObj.projColsF, evalOpsObj.numProjColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.avgCols, nextActualOpsObj.avgCols, evalOpsObj.avgColsP,
                        evalOpsObj.avgColsR, evalOpsObj.avgColsF, evalOpsObj.numAvgColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.minCols, nextActualOpsObj.minCols, evalOpsObj.minColsP,
                        evalOpsObj.minColsR, evalOpsObj.minColsF, evalOpsObj.numMinColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.maxCols, nextActualOpsObj.maxCols, evalOpsObj.maxColsP,
                        evalOpsObj.maxColsR, evalOpsObj.maxColsF, evalOpsObj.numMaxColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.sumCols, nextActualOpsObj.sumCols, evalOpsObj.sumColsP,
                        evalOpsObj.sumColsR, evalOpsObj.sumColsF, evalOpsObj.numSumColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.countCols, nextActualOpsObj.countCols, evalOpsObj.countColsP,
                        evalOpsObj.countColsR, evalOpsObj.countColsF, evalOpsObj.numCountColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.selCols, nextActualOpsObj.selCols, evalOpsObj.selColsP,
                        evalOpsObj.selColsR, evalOpsObj.selColsF, evalOpsObj.numSelColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.groupByCols, nextActualOpsObj.groupByCols, evalOpsObj.groupByColsP,
                        evalOpsObj.groupByColsR, evalOpsObj.groupByColsF, evalOpsObj.numGroupByColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.orderByCols, nextActualOpsObj.orderByCols, evalOpsObj.orderByColsP,
                        evalOpsObj.orderByColsR, evalOpsObj.orderByColsF, evalOpsObj.numOrderByColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.havingCols, nextActualOpsObj.havingCols, evalOpsObj.havingColsP,
                        evalOpsObj.havingColsR, evalOpsObj.havingColsF, evalOpsObj.numHavingColsQueries, evalOpsObj)
    compUpdateOpMetrics(predOpsObj.joinPreds, nextActualOpsObj.joinPreds, evalOpsObj.joinPredsP,
                        evalOpsObj.joinPredsR, evalOpsObj.joinPredsF, evalOpsObj.numJoinPredsColsQueries, evalOpsObj)
    compUpdateCondSelMetrics(predOpsObj, nextActualOpsObj, evalOpsObj)
    compUpdateCondJoinMetrics(predOpsObj, nextActualOpsObj, evalOpsObj)
    return

            
def createEvalMetricsOpWise(evalOpsObj):
    prevEpisode = -1
    rank = float("-inf")
    nextActualOpsObj = None
    predOpsObj = None
    with open(evalOpsObj.logFile) as f:
        for line in f:
            if line.startswith("#Episodes"):
                evalOpsObj.curEpisode = int(line.strip().split(";")[0].split(":")[1])
                numTokens = len(line.strip().split(";"))
                rank = int(line.strip().split(";")[numTokens-3].split(":")[1])
                if rank == -1: # this can happen when all predicted queries are equally bad
                    rank = 0
                assert rank >= 0 and rank < int(evalOpsObj.configDict['TOP_K'])
                MRR = float(1.0) / float(rank+1)
                if evalOpsObj.curEpisode != prevEpisode:
                    evalOpsObj.numEpQueries[evalOpsObj.curEpisode] = 1
                    assert evalOpsObj.curEpisode not in evalOpsObj.meanReciprocalRank
                    evalOpsObj.meanReciprocalRank[evalOpsObj.curEpisode] = MRR
                else:
                    evalOpsObj.numEpQueries[evalOpsObj.curEpisode] += 1
                    evalOpsObj.meanReciprocalRank[evalOpsObj.curEpisode] = (evalOpsObj.meanReciprocalRank[evalOpsObj.curEpisode] + MRR)
            elif line.startswith("Actual SQL"):
                evalOpsObj.curQueryIndex = -100
                nextActualOpsObj = nextActualOps()
            elif line.startswith("Predicted SQL Ops"):
                substrTokens = line.strip().split(":")[0].split(" ")
                evalOpsObj.curQueryIndex = int(substrTokens[len(substrTokens)-1])
                if evalOpsObj.curQueryIndex == rank:
                    predOpsObj = nextActualOps()
            elif line.startswith("---") and predOpsObj is not None and evalOpsObj is not None:
                computeF1(evalOpsObj, predOpsObj, nextActualOpsObj)
            elif evalOpsObj.curQueryIndex == -100:
                parseLineAddOp(line, nextActualOpsObj)
            elif evalOpsObj.curQueryIndex == rank:
                parseLineAddOp(line, predOpsObj)
            prevEpisode = evalOpsObj.curEpisode
    return evalOpsObj

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("-config", help="config file to parse", type=str, required=True)
    parser.add_argument("-log", help="log filename to analyze", type=str, required=True)
    #parser.add_argument("-lineNum", help="line Number to analyze", type=int, required=True)
    args = parser.parse_args()
    evalOpsObj = evalOps(args.config, args.log)
    evalOpsObj = createEvalMetricsOpWise(evalOpsObj)
    plotEvalMetricsOpWise(evalOpsObj)